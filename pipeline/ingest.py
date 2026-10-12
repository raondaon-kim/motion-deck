"""sources/ (md, txt, pdf, pptx, docx — mixed) -> one plain-text source document.
Deterministic text extraction; the storyboard LLM does the rewriting."""
from pathlib import Path


def _pdf(p):
    import fitz  # PyMuPDF
    with fitz.open(p) as d:
        return "\n\n".join(f"[page {i + 1}]\n{pg.get_text('text').strip()}" for i, pg in enumerate(d))


def _pptx(p):
    from pptx import Presentation
    out = []
    for i, sl in enumerate(Presentation(p).slides, 1):
        txt = [sh.text_frame.text.strip() for sh in sl.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text.strip()]
        notes = sl.notes_slide.notes_text_frame.text.strip() if sl.has_notes_slide and sl.notes_slide.notes_text_frame else ""
        out.append(f"[slide {i}]\n" + "\n".join(txt) + (f"\n(notes) {notes}" if notes else ""))
    return "\n\n".join(out)


def _docx(p):
    import docx
    return "\n".join(par.text for par in docx.Document(p).paragraphs)


READERS = {".md": lambda p: p.read_text(encoding="utf-8"), ".txt": lambda p: p.read_text(encoding="utf-8"),
           ".pdf": _pdf, ".pptx": _pptx, ".docx": _docx}


def ingest(src_dir):
    files = sorted(f for f in Path(src_dir).iterdir() if f.suffix.lower() in READERS)
    if not files: raise FileNotFoundError(f"no readable sources in {src_dir} (md, txt, pdf, pptx, docx)")
    parts = [f"# source: {f.name}\n\n{READERS[f.suffix.lower()](f).strip()}" for f in files]
    return "\n\n---\n\n".join(parts), [f.name for f in files]
